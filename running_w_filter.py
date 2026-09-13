import pandas as pd
import streamlit as st
import re
import io
from openpyxl.styles import PatternFill


# 全国车牌省份简称（31 个省级行政区）+ 车牌正则：
# 首字为省份简称；第二位为发牌机关代号（A-Z，不含 I/O）；
# 后部 5 位为普通车牌、6 位为新能源车牌（新能源第三位不限于 D/F，
# 各地使用 A/B/C/E/G/H/J/K 等扩展字母），字母均不含 I/O
PLATE_PROVINCES = '京津沪渝冀豫云辽黑湘皖鲁新苏浙赣鄂桂甘晋蒙陕吉闽贵粤青藏川宁琼'
PLATE_RE = re.compile(r'^[%s][A-HJ-NP-Z][A-HJ-NP-Z0-9]{5,6}$' % PLATE_PROVINCES)


# Normalize cell value to a clean string before regex matching:
# - Excel numeric cells become floats (13812345678.0) -> strip the ".0"
# - strip surrounding whitespace
def _norm(value):
    if pd.isna(value):
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# Function to validate data based on selected validators
def validate_data(df, validators):
    df = df.copy()

    # Apply validators to selected columns
    for column, validator in validators.items():
        if validator == 'Custom':
            custom_pattern = st.session_state.get(f'custom_validator_{column}', '')
            if custom_pattern:
                df[f'Valid_{column}'] = df[column].apply(
                    lambda x: bool(re.match(custom_pattern, _norm(x))))
            else:
                # No pattern provided: treat all values as valid
                df[f'Valid_{column}'] = True
        elif validator == 'Email':
            df[f'Valid_{column}'] = df[column].apply(
                lambda x: bool(re.match(r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$', _norm(x))))
        elif validator == 'EID':
            df[f'Valid_{column}'] = df[column].apply(
                lambda x: bool(re.match(r'^[0-9]+-[0-9]+-[0-9]+-\d$', _norm(x))))
        elif validator == 'Mobile No':
            df[f'Valid_{column}'] = df[column].apply(
                lambda x: bool(re.match(r'^1[3-9]\d{9}$', _norm(x))))
        elif validator == 'Plate':
            df[f'Valid_{column}'] = df[column].apply(
                lambda x: bool(PLATE_RE.match(_norm(x).upper())))
        elif validator == 'Date':
            df[f'Valid_{column}'] = pd.to_datetime(df[column], errors='coerce').notnull()
        elif validator == 'DateTime':
            df[f'Valid_{column}'] = pd.to_datetime(df[column], errors='coerce').notnull()
        elif validator is None:  # Handle columns with None validator
            df[f'Valid_{column}'] = True  # All values are considered valid
        # Add more validators as needed

    return df


# Build a marked Excel file: identical layout to the original file,
# only the invalid cells are highlighted in red.
def export_marked_file(df_original, df_validated, selected_columns):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        df_original.to_excel(writer, sheet_name='校验结果', index=False)
        ws = writer.book['校验结果']
        red = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
        col_index = {name: i + 1 for i, name in enumerate(df_original.columns)}
        for col in selected_columns:
            invalid = ~df_validated[f'Valid_{col}'].astype(bool)
            for row_pos, flag in enumerate(invalid):
                if flag:
                    ws.cell(row=row_pos + 2, column=col_index[col]).fill = red
    buf.seek(0)
    return buf


# Streamlit UI
st.title('Data Validator')

# File upload
uploaded_file = st.file_uploader('Upload CSV or Excel file', type=['csv', 'xls', 'xlsx'])

if uploaded_file is not None:
    # Read the uploaded file
    try:
        df = pd.read_csv(uploaded_file) if uploaded_file.name.endswith('.csv') else pd.read_excel(uploaded_file)
    except Exception as e:
        st.error(f'Error reading file: {str(e)}')
    else:
        # Display DataFrame
        st.write('Original DataFrame:')
        st.write(df)

        # Select columns to work with
        selected_columns = st.multiselect('Select columns to work with', df.columns)

        # Allow user to add custom validator
        custom_validators = {}
        for column in selected_columns:
            custom_pattern = st.text_input(f'Enter custom regex pattern for column "{column}"')
            if custom_pattern:
                custom_validators[column] = custom_pattern

        # Define validators
        validators = {
            'None': None,
            'Email': 'Email',
            'EID': 'EID',
            'Mobile No (11位手机号)': 'Mobile No',
            'License Plate (车牌号)': 'Plate',
            'Date': 'Date',
            'DateTime': 'DateTime',
            'Custom (自定义)': 'Custom',
            # Add more validators here
        }

        # Assign validator for each column
        selected_validators = {}
        for column in selected_columns:
            validator = st.selectbox(f'Select validator for column "{column}"', list(validators.keys()))
            selected_validators[column] = validators[validator]

            # Store custom validator pattern in session state
            if validator == 'Custom':
                st.session_state[f'custom_validator_{column}'] = custom_validators.get(column, '')

        if st.button('Run Validator'):
            # Validate the data
            df_validated = validate_data(df[selected_columns], selected_validators)

            # Display validation summary
            validation_results = pd.DataFrame()
            for column, validator in selected_validators.items():
                validation_results.loc['Validation Score', column] = df_validated[f'Valid_{column}'].mean() * 100
                validation_results.loc['Completeness Score', column] = df_validated[column].notnull().mean() * 100

            st.write('Validation Summary:')
            st.write(validation_results)

            # Display non-valid rows
            valid_columns = [col for col in df_validated.columns if col.startswith('Valid_')]
            non_valid_rows = df_validated[~df_validated[valid_columns].all(axis=1)] if valid_columns else df_validated.iloc[0:0]

            if not non_valid_rows.empty:
                st.write('Filtered Dataset (Non-Valid Rows Only):')
                st.write(non_valid_rows)
            else:
                st.write('No non-valid rows found.')

            # Download the marked file (invalid cells highlighted in red)
            if selected_validators:
                st.write('下载标记结果：')
                marked_file = export_marked_file(df, df_validated, list(selected_validators.keys()))
                st.download_button(
                    '下载标记后的 Excel（无效单元格标红）',
                    marked_file,
                    'validation_marked.xlsx',
                    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                )
